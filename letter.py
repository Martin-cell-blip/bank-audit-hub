"""
询证函生成：银行存款/借款 → 银行询证函（一封覆盖该行全部事项，存款+借款分行列示）；
应收/应付/其他应收款 → 企业往来询证函。
输出两种形态：可打印 HTML（浏览器直接转 PDF）与可下载 Word (.docx)。
模板参照中注协《银行函证及回函工作操作指引》与往来函标准格式，做了教学化简化。
"""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

import db

FIRM = "天津信诚会计师事务所（特殊普通合伙）"
FY = db.FY
BALANCE = f"{FY}-12-31"
BANK_ACCOUNTS = set(db.BANK_ACCOUNTS)

# 往来函必备免责条款：防止被询证方误当催款函而拒回
TRADE_DISCLAIMER = "本函仅为复核账目之用，并非催款结算依据。"


def _fmt(n) -> str:
    try:
        return f"{float(n):,.2f}"
    except Exception:
        return "0.00"


def is_bank(conf: dict) -> bool:
    return conf.get("account") in BANK_ACCOUNTS


def _today() -> str:
    return date.today().strftime("%Y 年 %m 月 %d 日")


# --------------------------------------------------------------------------- #
# HTML（可打印）
# --------------------------------------------------------------------------- #
def letter_html(conf: dict, eng: dict, rows: list | None = None) -> str:
    """rows: 同一封函覆盖的全部台账/函证行（银行函证聚合用）；None → 仅 conf 一行。"""
    rows = rows or [conf]
    entity = eng.get("entity", "")
    no = conf.get("confirm_no", "")
    method = conf.get("method", "积极式")
    bank = conf.get("bank_name") or conf.get("counterparty", "")

    if is_bank(conf):
        addressee = f"致：{bank}"
        intro = (f"本公司聘请 {FIRM} 对本公司 <b>{FY} 年度</b> 财务报表进行审计。按照中国注册会计师"
                 f"审计准则的要求，应当对本公司与贵行相关的信息予以函证。下列各项数据出自本公司账簿记录，"
                 f"如与贵行记录相符，请在本函下端“信息证明无误”处签章证明；如有不符，请在“信息不符”处"
                 f"列明不符项目及具体金额。<b>本函包含本公司在贵行的全部存款、借款及其他事项，请一并核对。</b>")
        head = "<tr><th>项目</th><th>科目</th><th>账面余额（人民币元）</th><th>截止日</th></tr>"
        body_rows = "".join(
            f"<tr><td>{r.get('counterparty','')}</td><td>{r.get('account','')}</td>"
            f"<td class='num'>{_fmt(r.get('book_amount'))}</td><td>{BALANCE}</td></tr>"
            for r in rows)
        sign_yes = "上述各项信息经核对无误。"
        sign_no = "上述信息与本行记录不符，不符项目及金额如下："
        disclaimer = ""
    else:
        cp = conf.get("counterparty", "")
        rel = "应收/其他应收" if conf.get("account") in ("应收账款", "其他应收款") else "应付"
        addressee = f"致：{cp}"
        neg = ("如与贵单位记录不符，请于收到本函后回函说明；如相符，可不必回函（消极式）。"
               if method == "消极式" else "如相符请签章确认；如不符请列明差异（积极式）。")
        intro = (f"本公司聘请 {FIRM} 对本公司 <b>{FY} 年度</b> 财务报表进行审计。现将本公司与贵单位"
                 f"截至 {BALANCE} 的往来款项余额列示如下，请核对。{neg}")
        head = "<tr><th>往来单位</th><th>科目</th><th>本公司账面余额（人民币元）</th><th>截止日</th></tr>"
        body_rows = (f"<tr><td>{cp}</td><td>{conf.get('account','')}（{rel}）</td>"
                     f"<td class='num'>{_fmt(conf.get('book_amount'))}</td><td>{BALANCE}</td></tr>")
        sign_yes = "上述往来余额与本单位记录相符。"
        sign_no = "上述往来余额与本单位记录不符，差异如下："
        disclaimer = f"<p class='disc'>{TRADE_DISCLAIMER}</p>"

    # 回函登记（逐行）
    reply_rows = "".join(
        f"<div>· {r.get('account','')} {r.get('counterparty','')}：回函 <b>{_fmt(r.get('reply_amount'))}</b> 元，"
        f"差异 <b class=\"{'diff' if r.get('difference') else ''}\">{_fmt(r.get('difference'))}</b> 元"
        f"{('（' + r.get('diff_reason') + '）') if r.get('diff_reason') else ''}</div>"
        for r in rows if r.get("reply_amount") is not None)
    reply_block = f"<div class='reply'><div class='reply-h'>【回函登记】</div>{reply_rows}</div>" if reply_rows else ""

    return f"""<!doctype html><html lang="zh"><head><meta charset="utf-8">
<title>询证函 {no}</title>
<style>
  @page {{ size: A4; margin: 22mm; }}
  body {{ font-family: "SimSun","宋体",serif; color:#111; line-height:1.9; font-size:15px; }}
  .wrap {{ max-width: 720px; margin: 0 auto; }}
  h1 {{ text-align:center; font-size:24px; letter-spacing:8px; margin-bottom:2px; }}
  .sub {{ text-align:center; color:#555; margin-bottom:18px; }}
  .meta {{ display:flex; justify-content:space-between; font-size:13px; color:#333; border-bottom:1px solid #999; padding-bottom:6px; }}
  .addr {{ font-weight:bold; margin:16px 0 6px; }}
  .tb {{ width:100%; border-collapse:collapse; margin:12px 0; }}
  .tb th, .tb td {{ border:1px solid #333; padding:7px 9px; font-size:14px; }}
  .tb th {{ background:#f0f0f0; }}
  .num {{ text-align:right; font-variant-numeric: tabular-nums; }}
  .disc {{ font-size:13px; color:#b00; margin:6px 0; }}
  .sign {{ margin-top:26px; }}
  .sign .row {{ display:flex; gap:40px; margin:14px 0; }}
  .box {{ flex:1; border:1px solid #333; padding:10px 12px; min-height:70px; }}
  .box h4 {{ margin:0 0 6px; font-size:14px; }}
  .foot {{ margin-top:30px; text-align:right; }}
  .reply {{ margin-top:18px; padding:10px 12px; background:#fbfaf3; border:1px dashed #b8860b; font-size:14px; }}
  .reply-h {{ font-weight:bold; margin-bottom:4px; }}
  .diff {{ color:#c0392b; }}
  @media print {{ .noprint {{ display:none; }} }}
  .noprint {{ text-align:center; margin:14px 0; }}
  .btn {{ padding:8px 16px; border:1px solid #333; background:#fff; cursor:pointer; border-radius:6px; }}
</style></head>
<body><div class="wrap">
  <div class="noprint"><button class="btn" onclick="window.print()">打印 / 另存为 PDF</button></div>
  <h1>询 证 函</h1>
  <div class="sub">{'银行询证函' if is_bank(conf) else '企业往来询证函'}（{method}）</div>
  <div class="meta"><span>编号：{no}</span><span>被审计单位：{entity}</span></div>
  <div class="addr">{addressee}：</div>
  <p>{intro}</p>
  <table class="tb">{head}{body_rows}</table>
  {disclaimer}
  {reply_block}
  <div class="sign">
    <div class="row">
      <div class="box"><h4>信息证明无误</h4>{sign_yes}<br><br>（签章）　　经办人：　　　　日期：</div>
      <div class="box"><h4>信息不符</h4>{sign_no}<br><br>（签章）　　经办人：　　　　日期：</div>
    </div>
  </div>
  <div class="foot">
    被审计单位（盖章）：{entity}<br>
    会计师事务所：{FIRM}<br>
    出函日期：{_today()}
  </div>
</div></body></html>"""


# --------------------------------------------------------------------------- #
# Word (.docx)
# --------------------------------------------------------------------------- #
def letter_docx(conf: dict, eng: dict, rows: list | None = None) -> bytes:
    rows = rows or [conf]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    def p(text="", align=None, bold=False, size=None):
        par = doc.add_paragraph()
        run = par.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align:
            par.alignment = align
        return par

    p("询 证 函", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=20)
    p(("银行询证函" if is_bank(conf) else "企业往来询证函") + f"（{conf.get('method','积极式')}）",
      align=WD_ALIGN_PARAGRAPH.CENTER)
    p(f"编号：{conf.get('confirm_no','')}    被审计单位：{eng.get('entity','')}")

    if is_bank(conf):
        recv = conf.get("bank_name") or conf.get("counterparty", "")
        p(f"致：{recv}：", bold=True)
        p(f"本公司聘请 {FIRM} 对本公司 {FY} 年度财务报表进行审计。下列各项数据出自本公司账簿记录，"
          f"如与贵行记录相符，请签章证明；如有不符，请列明不符项目及金额。本函包含本公司在贵行的"
          f"全部存款、借款及其他事项，请一并核对。")
        cols = ("往来单位/开户行", "科目", "账面余额(元)", "截止日")
    else:
        recv = conf.get("counterparty", "")
        p(f"致：{recv}：", bold=True)
        p(f"本公司聘请 {FIRM} 对本公司 {FY} 年度财务报表进行审计。现将本公司与贵单位"
          f"截至 {BALANCE} 的往来款项余额列示如下，请核对确认。")
        p(TRADE_DISCLAIMER, bold=True)
        cols = ("往来单位", "科目", "账面余额(元)", "截止日")
        rows = [conf]

    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, name in enumerate(cols):
        t.rows[0].cells[i].text = name
    for r in rows:
        cells = t.add_row().cells
        cells[0].text = r.get("counterparty", "") or recv
        cells[1].text = r.get("account", "")
        cells[2].text = _fmt(r.get("book_amount"))
        cells[3].text = BALANCE

    replied = [r for r in rows if r.get("reply_amount") is not None]
    if replied:
        p()
        for r in replied:
            rr = f"【回函登记】{r.get('account','')} {r.get('counterparty','')}：回函 {_fmt(r.get('reply_amount'))} 元；" \
                 f"差异 {_fmt(r.get('difference'))} 元"
            if r.get("diff_reason"):
                rr += f"；差异归因：{r.get('diff_reason')}"
            p(rr)

    p()
    p("信息证明无误（签章）：________________    经办人：________  日期：________")
    p("信息不符，差异如下（签章）：____________________________________________")
    p()
    p(f"被审计单位（盖章）：{eng.get('entity','')}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    p(f"会计师事务所：{FIRM}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    p(f"出函日期：{_today()}", align=WD_ALIGN_PARAGRAPH.RIGHT)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

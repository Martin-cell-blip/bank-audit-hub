"""
银行审计「函证 + 多线任务」协同平台 —— FastAPI 后端。
运行：  uvicorn app:app --reload  然后打开 http://127.0.0.1:8000
"""
from __future__ import annotations

import os
from contextlib import asynccontextmanager
from datetime import date
from decimal import Decimal
from urllib.parse import quote

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import HTMLResponse, Response, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import db
import letter
import workflow


@asynccontextmanager
async def lifespan(_app: FastAPI):
    db.ensure_seeded()
    yield
    db.close()


app = FastAPI(title="银行审计函证与多线任务协同平台", lifespan=lifespan)
HERE = os.path.dirname(__file__)


# --------------------------------------------------------------------------- #
# 总览 Dashboard
# --------------------------------------------------------------------------- #
@app.get("/api/overview")
def overview():
    counts = {r["status"]: r["n"] for r in db.query(
        "SELECT status, COUNT(*) n FROM confirmations GROUP BY status")}
    for s in db.STATUSES:
        counts.setdefault(s, 0)
    total = sum(counts.values())
    replied = counts["已回函"] + counts["差异待跟进"] + counts["待复核"] + counts["已结项"]

    agg = db.query("""
        SELECT
          SUM(CASE WHEN difference IS NOT NULL AND difference <> 0 THEN 1 ELSE 0 END) AS diff_cnt,
          COALESCE(SUM(CASE WHEN difference IS NOT NULL AND difference <> 0 THEN ABS(difference) ELSE 0 END),0) AS diff_amt,
          COALESCE(SUM(book_amount),0) AS book_total
        FROM confirmations""")[0]

    overdue = db.query("""
        SELECT COUNT(*) n FROM engagements
        WHERE due_date < CURRENT_DATE AND stage <> '报告'""")[0]["n"]

    # 账龄分布：仅往来款项（应收/应付/其他应收）——银行存款/借款无账龄概念
    recv = "(" + ",".join("'%s'" % a for a in db.RECEIVABLE_ACCOUNTS) + ")"
    aging = db.query(f"""
        WITH b AS (
          SELECT CASE
            WHEN CAST(julianday(as_of)-julianday(entry_date) AS INTEGER) <= 30 THEN '0-30天'
            WHEN CAST(julianday(as_of)-julianday(entry_date) AS INTEGER) <= 90 THEN '31-90天'
            WHEN CAST(julianday(as_of)-julianday(entry_date) AS INTEGER) <= 180 THEN '91-180天'
            WHEN CAST(julianday(as_of)-julianday(entry_date) AS INTEGER) <= 365 THEN '181-365天'
            ELSE '1年以上' END AS bucket,
          book_amount
          FROM ledger WHERE account IN {recv})
        SELECT bucket, COUNT(*) cnt, SUM(book_amount) amt FROM b GROUP BY bucket""")
    order = ['0-30天', '31-90天', '91-180天', '181-365天', '1年以上']
    aging.sort(key=lambda x: order.index(x["bucket"]) if x["bucket"] in order else 99)

    # 函证覆盖率（金额口径，按科目）——审计质量控制核心指标
    coverage = db.query("""
        SELECT l.account,
          SUM(l.book_amount) AS total_amt,
          COALESCE(SUM(CASE WHEN EXISTS(
            SELECT 1 FROM confirmations c WHERE c.ledger_id = l.id
          ) THEN l.book_amount ELSE 0 END), 0) AS confirmed_amt
        FROM ledger l GROUP BY l.account""")
    for cov in coverage:
        cov["rate"] = round(cov["confirmed_amt"] / cov["total_amt"] * 100, 2) if cov["total_amt"] else 0
        # 银行存款监管红线：应 100% 函证（含零余额户、销户）
        cov["alert"] = (cov["account"] == "银行存款" and cov["rate"] < 100)
    coverage.sort(key=lambda x: -x["total_amt"])

    entity_progress = db.query("""
        SELECT e.entity,
               COUNT(c.id) AS total,
               SUM(CASE WHEN c.status IN ('已回函','差异待跟进','待复核','已结项') THEN 1 ELSE 0 END) AS replied
        FROM engagements e LEFT JOIN confirmations c ON c.engagement_id = e.id
        GROUP BY e.entity ORDER BY e.entity""")

    return {
        "status_counts": counts,
        "total": total,
        "replied": replied,
        "reply_rate": round(replied / total * 100, 1) if total else 0,
        "diff_cnt": agg["diff_cnt"],
        "diff_amt": agg["diff_amt"],
        "book_total": agg["book_total"],
        "overdue": overdue,
        "aging": aging,
        "coverage": coverage,
        "entity_progress": entity_progress,
        "fy": db.FY,
        "as_of": db.REPORT_DATE.isoformat(),
        "today": date.today().isoformat(),
    }


# --------------------------------------------------------------------------- #
# 函证看板
# --------------------------------------------------------------------------- #
URGE_DAYS = 30   # 发函后超过 N 天未回 → 应催函 / 转替代程序


@app.get("/api/confirmations")
def list_confirmations():
    rows = db.query("""
        SELECT c.*, e.entity, e.audit_area,
          p.name AS preparer_name, r.name AS reviewer_name,
          (SELECT COUNT(*) FROM evidence_files ef WHERE ef.confirmation_id=c.id) AS evidence_count,
          (SELECT COUNT(*) FROM review_decisions rd WHERE rd.confirmation_id=c.id) AS review_count,
          CASE WHEN c.sent_date IS NOT NULL
               THEN CAST(julianday(CURRENT_DATE)-julianday(c.sent_date) AS INTEGER) END AS days_since_sent
        FROM confirmations c JOIN engagements e ON e.id = c.engagement_id
        JOIN users p ON p.id=c.preparer_id
        JOIN users r ON r.id=c.reviewer_id
        ORDER BY c.confirm_no""")
    for r in rows:
        d = r.get("days_since_sent")
        # 已发出且超期未回：积极式须催函/转替代程序；银行存款不得以替代程序替代
        r["needs_followup"] = bool(r["status"] == "已发出" and d is not None and d >= URGE_DAYS)
        r["no_alt_procedure"] = r["account"] in db.BANK_ACCOUNTS
    return rows


class StatusIn(BaseModel):
    status: str
    actor_id: int = 1


class ActionIn(BaseModel):
    action: str
    actor_id: int = 1
    reply_amount: float | None = None
    conclusion: str | None = None
    decision: str | None = None
    notes: str | None = None


class ReasonIn(BaseModel):
    reason: str
    actor_id: int = 1


def _get_conf(cid: int) -> dict:
    try:
        return workflow.get_confirmation(cid)
    except workflow.WorkflowError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc


@app.post("/api/confirmations/{cid}/status")
def set_status(cid: int, body: StatusIn):
    conf = _get_conf(cid)
    if body.status == conf["status"]:
        return conf
    action = "send" if body.status == "已发出" else "reset" if body.status == "待发函" else None
    if action is None:
        raise HTTPException(409, "该状态必须通过回函、提交复核或复核决定按钮推进")
    try:
        return workflow.apply_action(cid, action, body.actor_id)
    except workflow.WorkflowError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc


@app.post("/api/confirmations/{cid}/action")
def act(cid: int, body: ActionIn):
    try:
        return workflow.apply_action(
            cid,
            body.action,
            body.actor_id,
            reply_amount=body.reply_amount,
            conclusion=body.conclusion,
            decision=body.decision,
            notes=body.notes,
        )
    except workflow.WorkflowError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc


@app.post("/api/confirmations/{cid}/reason")
def set_reason(cid: int, body: ReasonIn):
    try:
        return workflow.set_reason(cid, body.actor_id, body.reason)
    except workflow.WorkflowError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc


@app.get("/api/users")
def users():
    return db.query("SELECT id, name, role FROM users ORDER BY id")


@app.get("/api/confirmations/{cid}/evidence")
def list_evidence(cid: int):
    _get_conf(cid)
    return db.query(
        """SELECT e.id, e.evidence_type, e.original_name, e.sha256, e.uploaded_at,
                  e.version, u.name AS uploader_name
           FROM evidence_files e JOIN users u ON u.id=e.uploaded_by
           WHERE e.confirmation_id=? ORDER BY e.id""",
        [cid],
    )


@app.post("/api/confirmations/{cid}/evidence")
async def upload_evidence(
    cid: int,
    file: UploadFile = File(...),
    evidence_type: str = Form("回函扫描件"),
    actor_id: int = Form(1),
):
    try:
        return workflow.add_evidence(
            cid,
            actor_id,
            evidence_type,
            file.filename or "evidence.bin",
            await file.read(),
        )
    except workflow.WorkflowError as exc:
        raise HTTPException(exc.status_code, str(exc)) from exc


@app.get("/api/evidence/{evidence_id}/download")
def download_evidence(evidence_id: int):
    rows = db.query("SELECT * FROM evidence_files WHERE id=?", [evidence_id])
    if not rows:
        raise HTTPException(404, "证据不存在")
    evidence = rows[0]
    path = os.path.abspath(evidence["stored_path"])
    evidence_root = os.path.abspath(db.EVIDENCE_DIR)
    if os.path.commonpath([path, evidence_root]) != evidence_root or not os.path.isfile(path):
        raise HTTPException(404, "证据文件不可用")
    return FileResponse(path, filename=evidence["original_name"])


@app.get("/api/confirmations/{cid}/letter")
def get_letter(cid: int, fmt: str = "html"):
    conf = _get_conf(cid)
    eng = db.query("SELECT * FROM engagements WHERE id=?", [conf["engagement_id"]])[0]
    # 银行询证函：一封覆盖该行（同一被审计单位）全部存款/借款事项 → 聚合同 bank_name+entity 的兄弟函证
    rows = None
    if conf["account"] in db.BANK_ACCOUNTS and conf["bank_name"]:
        rows = db.query("""
            SELECT c.* FROM confirmations c JOIN engagements e ON e.id = c.engagement_id
            WHERE c.bank_name = ? AND e.entity = ? ORDER BY c.account, c.confirm_no""",
            [conf["bank_name"], eng["entity"]])
    if fmt == "docx":
        data = letter.letter_docx(conf, eng, rows)
        no = conf["confirm_no"]
        fname = quote(f"询证函_{no}.docx")          # RFC 5987：HTTP 头仅支持 latin-1，中文名须 URL 编码
        disp = f"attachment; filename=\"letter_{no}.docx\"; filename*=UTF-8''{fname}"
        return Response(
            content=data,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": disp},
        )
    return HTMLResponse(letter.letter_html(conf, eng, rows))


# --------------------------------------------------------------------------- #
# 台账 + 批量生成
# --------------------------------------------------------------------------- #
@app.get("/api/ledger")
def list_ledger():
    return db.query("""
        SELECT l.*, e.entity,
          CAST(julianday(l.as_of)-julianday(l.entry_date) AS INTEGER) AS age_days,
          EXISTS(SELECT 1 FROM confirmations c WHERE c.ledger_id = l.id) AS has_conf
        FROM ledger l JOIN engagements e ON e.id = l.engagement_id
        ORDER BY l.id""")


_AREA_CODE = {"银行存款": "YH", "应收账款": "YS", "应付账款": "YF",
              "短期借款": "JK", "长期借款": "JK", "其他应收款": "QT"}


def _new_confirm_no(area: str) -> str:
    code = _AREA_CODE.get(area, "XX")
    n = db.query("SELECT COUNT(*) n FROM confirmations WHERE confirm_no LIKE ?",
                 [f"YQ{db.FY}-{code}-%"])[0]["n"]
    return f"YQ{db.FY}-{code}-{n + 1:03d}"


@app.post("/api/ledger/generate")
def generate():
    rows = db.query("""
        SELECT l.* FROM ledger l
        WHERE NOT EXISTS(SELECT 1 FROM confirmations c WHERE c.ledger_id = l.id)""")
    created = []
    for r in rows:
        cid = db.next_id("confirmations")
        no = _new_confirm_no(r["account"])
        method = db.pick_method(r["account"], r["book_amount"])
        db.execute(
            """INSERT INTO confirmations
               (id,ledger_id,engagement_id,confirm_no,counterparty,bank_name,account,method,
                book_amount,reply_amount,difference,status,sent_date,reply_date,updated_at,
                diff_reason,preparer_id,reviewer_id,preparer_conclusion)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            [
                cid, r["id"], r["engagement_id"], no, r["counterparty"], r["bank_name"],
                r["account"], method, float(r["book_amount"]), None, None, "待发函",
                None, None, date.today().isoformat(), None, 1, 2, None,
            ],
        )
        db.log_event(cid, 1, "生成函证", f"{no} · {r['counterparty']}", None, "待发函")
        created.append(no)
    return {"created": len(created), "nos": created}


@app.post("/api/ledger/import")
async def import_ledger(file: UploadFile = File(...), engagement_id: int = Form(...)):
    """导入 Excel 台账（列：往来单位/开户行, 科目, 账面余额, 入账日[可选]）。"""
    import openpyxl
    import io as _io
    data = await file.read()
    wb = openpyxl.load_workbook(_io.BytesIO(data), data_only=True)
    ws = wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    ci_cp = col("往来单位", "开户行", "往来单位/开户行", "单位名称")
    ci_ac = col("科目", "会计科目")
    ci_amt = col("账面余额", "余额", "金额", "账面金额")
    ci_dt = col("入账日", "入账日期", "日期")
    if ci_cp is None or ci_amt is None:
        raise HTTPException(400, "缺少必需列：往来单位/开户行、账面余额")

    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue
        cp = str(row[ci_cp]).strip() if row[ci_cp] is not None else ""
        if not cp:
            continue
        account = str(row[ci_ac]).strip() if ci_ac is not None and row[ci_ac] else "应收账款"
        try:
            amt = Decimal(str(row[ci_amt]))
        except Exception:
            continue
        entry = row[ci_dt] if ci_dt is not None else None
        try:
            entry_dt = entry.date() if hasattr(entry, "date") else (
                date.fromisoformat(str(entry)[:10]) if entry else db.REPORT_DATE)
        except Exception:
            entry_dt = db.REPORT_DATE
        bank = cp if account in letter.BANK_ACCOUNTS else ""
        lid = db.next_id("ledger")
        db.execute("INSERT INTO ledger VALUES (?,?,?,?,?,?,?,?,?)",
                   [lid, engagement_id, cp, bank, account, float(amt), "CNY",
                    entry_dt.isoformat(), db.REPORT_DATE.isoformat()])
        n += 1
    db.log_event(None, 1, "导入台账", f"{file.filename}：{n} 行 → 项目#{engagement_id}")
    return {"imported": n}


# --------------------------------------------------------------------------- #
# 多线任务 / 审计项目
# --------------------------------------------------------------------------- #
@app.get("/api/engagements")
def list_engagements():
    return db.query("""
        SELECT e.*,
          CAST(julianday(e.due_date)-julianday(CURRENT_DATE) AS INTEGER) AS days_left,
          (e.due_date < CURRENT_DATE AND e.stage <> '报告') AS overdue,
          COUNT(c.id) AS conf_total,
          SUM(CASE WHEN c.status IN ('已回函','差异待跟进','待复核','已结项') THEN 1 ELSE 0 END) AS conf_replied,
          SUM(CASE WHEN c.status = '差异待跟进' THEN 1 ELSE 0 END) AS conf_diff
        FROM engagements e LEFT JOIN confirmations c ON c.engagement_id = e.id
        GROUP BY e.id, e.entity, e.audit_area, e.period, e.lead, e.stage, e.start_date, e.due_date
        ORDER BY e.id""")


class AdvanceIn(BaseModel):
    direction: str = "next"


@app.post("/api/engagements/{eid}/advance")
def advance(eid: int, body: AdvanceIn):
    r = db.query("SELECT * FROM engagements WHERE id=?", [eid])
    if not r:
        raise HTTPException(404, "项目不存在")
    cur = r[0]["stage"]
    i = db.STAGES.index(cur) if cur in db.STAGES else 0
    i = min(i + 1, len(db.STAGES) - 1) if body.direction == "next" else max(i - 1, 0)
    db.execute("UPDATE engagements SET stage=? WHERE id=?", [db.STAGES[i], eid])
    return db.query("SELECT * FROM engagements WHERE id=?", [eid])[0]


# --------------------------------------------------------------------------- #
# 杂项
# --------------------------------------------------------------------------- #
@app.get("/api/events")
def events(limit: int = 12):
    return db.query(
        """SELECT e.*, u.name AS actor_name FROM events e
           LEFT JOIN users u ON u.id=e.actor_id ORDER BY e.id DESC LIMIT ?""",
        [limit],
    )


@app.post("/api/reset")
def reset():
    return db.seed()


@app.get("/", response_class=HTMLResponse)
def index():
    return FileResponse(os.path.join(HERE, "static", "index.html"))


app.mount("/static", StaticFiles(directory=os.path.join(HERE, "static")), name="static")
